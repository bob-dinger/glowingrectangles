"""Write ~/Desktop/beatles_missing.xlsx: every Beatles song in parcels.songs that
is missing melody and/or chords, with whether a ready paste asset exists.

Re-run any time after a sync to get the current gap list.
"""
import os, re, glob, psycopg2
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
load_dotenv('/Users/robert/Desktop/themap/themap_claude/.env')

def norm(s): return re.sub(r'[^a-z0-9]', '', (s or '').lower())

def main():
    c = psycopg2.connect(host=os.environ['DB_HOST'], dbname=os.environ['DB_NAME'],
        user=os.environ['DB_USER'], password=os.environ['DB_PASSWORD'], port=os.environ.get('DB_PORT', 5432))
    cur = c.cursor()
    cur.execute("""select distinct on (norm) title, sec, mel, chd from (
        select title, lower(regexp_replace(title,'[^a-zA-Z0-9]','','g')) norm,
          coalesce(jsonb_array_length(hookpad_json->'sections'),0) sec,
          coalesce(jsonb_array_length(hookpad_json->'notes'),0) mel,
          coalesce(jsonb_array_length(hookpad_json->'chords'),0) chd
        from parcels.songs where slug like 'beatles_%' and hookpad_json is not null) t
        order by norm, mel desc, chd desc""")
    rows = cur.fetchall()

    PP = os.path.expanduser('~/Desktop/pollack_pastes')
    melodies = {norm(os.path.basename(p)[len('melody_'):-4]) for p in glob.glob(PP + '/melody_*.txt')}
    chords = {norm(os.path.basename(p)[:-4]) for p in glob.glob(PP + '/*.txt')
              if not os.path.basename(p).startswith(('melody_', '_melody'))}

    missing = []
    for title, sec, mel, chd in rows:
        lack = ([] if mel else ['Melody']) + ([] if chd else ['Chords'])
        if lack:
            n = norm(title)
            missing.append([title, ' + '.join(lack), sec, mel, chd,
                            'yes' if n in melodies else '', 'yes' if n in chords else ''])
    missing.sort(key=lambda r: (r[1], r[0].lower()))

    wb = Workbook(); ws = wb.active; ws.title = 'Beatles Missing'
    ws.append(['Song', 'Missing', 'Sections', 'Notes', 'Chords', 'Melody paste ready?', 'Chord paste ready?'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center')
    green = PatternFill('solid', fgColor='C6EFCE')
    for r in missing:
        ws.append(r)
        i = ws.max_row
        if r[5]: ws.cell(i, 6).fill = green
        if r[6]: ws.cell(i, 7).fill = green
    for i, w in enumerate([34, 16, 10, 8, 8, 20, 20], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    out = os.path.expanduser('~/Desktop/beatles_missing.xlsx')
    wb.save(out)
    print(f'wrote {out}  ({len(missing)} incomplete songs)')
    bym = {}
    for r in missing: bym[r[1]] = bym.get(r[1], 0) + 1
    print('breakdown:', bym)

if __name__ == '__main__':
    main()
