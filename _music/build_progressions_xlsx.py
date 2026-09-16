"""Build /Users/robert/Desktop/four_chord_progressions_by_song.xlsx from the CSV.

Coloring is by PITCH CLASS of the chord root in the home key (C major / A minor reference):
  - Diatonic positions (0,2,4,5,7,9,11) → solid color (red/orange/yellow/green/blue/purple/pink)
  - Chromatic positions (1,3,6,8,10)    → diagonal stripes of the two neighbors

This makes V/V color as orange (root pc=2) instead of blue, and bVII (= IV/IV) as
purple+pink diagonal stripes. Secondary dominants inherit the color of their actual
chord, not the Roman-letter prefix.
"""
import csv, re, sys
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

CSV_PATH = '/Users/robert/Desktop/four_chord_progressions_by_song.csv'
XLSX_PATH = '/Users/robert/Desktop/four_chord_progressions_by_song.xlsx'

# --- color palette: 7 diatonic colors + matching text color ----------------------
DEG_RGB = {
    1: 'E84545',  # I    red
    2: 'F0A040',  # II   orange
    3: 'E8C828',  # III  yellow
    4: '50C878',  # IV   green
    5: '5090F0',  # V    blue
    6: '7040B0',  # VI   purple
    7: 'E070B0',  # VII  pink
}
DEG_TEXT = {1: 'FFFFFF', 2: '000000', 3: '000000', 4: '000000',
            5: 'FFFFFF', 6: 'FFFFFF', 7: 'FFFFFF'}

# pc 0..11 → (pattern, fg_hex, bg_hex_or_None, text_hex)
# Chromatic cells use 'darkUp' (thick diagonal stripes). fg = stripe, bg = field.
PC_STYLE = {
    0:  ('solid',  DEG_RGB[1], None,        DEG_TEXT[1]),                  # I    red
    1:  ('darkUp', DEG_RGB[1], DEG_RGB[2],  'FFFFFF'),                     # bII  red + orange
    2:  ('solid',  DEG_RGB[2], None,        DEG_TEXT[2]),                  # II   orange
    3:  ('darkUp', DEG_RGB[2], DEG_RGB[3],  '000000'),                     # bIII orange + yellow
    4:  ('solid',  DEG_RGB[3], None,        DEG_TEXT[3]),                  # III  yellow
    5:  ('solid',  DEG_RGB[4], None,        DEG_TEXT[4]),                  # IV   green
    6:  ('darkUp', DEG_RGB[4], DEG_RGB[5],  'FFFFFF'),                     # bV   green + blue
    7:  ('solid',  DEG_RGB[5], None,        DEG_TEXT[5]),                  # V    blue
    8:  ('darkUp', DEG_RGB[5], DEG_RGB[6],  'FFFFFF'),                     # bVI  blue + purple
    9:  ('solid',  DEG_RGB[6], None,        DEG_TEXT[6]),                  # VI   purple
    10: ('darkUp', DEG_RGB[6], DEG_RGB[7],  'FFFFFF'),                     # bVII purple + pink
    11: ('solid',  DEG_RGB[7], None,        DEG_TEXT[7]),                  # VII  pink
}

ROMAN_VAL = {'I': 1, 'II': 2, 'III': 3, 'IV': 4, 'V': 5, 'VI': 6, 'VII': 7}
MAJOR_INTERVALS = {1: 0, 2: 2, 3: 4, 4: 5, 5: 7, 6: 9, 7: 11}
PC_TO_NAME = ['C', 'Db', 'D', 'Eb', 'E', 'F', 'Gb', 'G', 'Ab', 'A', 'Bb', 'B']


def parse_roman_part(s):
    """Parse a single Roman-numeral chord token (no slash) into pc + quality + leftover."""
    if not s:
        return None
    acc = 0
    if s.startswith('b'):
        acc = -1; s = s[1:]
    elif s.startswith('#'):
        acc = 1; s = s[1:]
    su = s.upper()
    for L in (3, 2, 1):
        c = su[:L]
        if c in ROMAN_VAL:
            return {
                'pc': (MAJOR_INTERVALS[ROMAN_VAL[c]] + acc) % 12,
                'is_minor': s[:L].islower(),
                'rest': s[L:],
            }
    return None


def token_to_pc(token):
    """Pitch class (0-11) of the chord root in the home key, or None."""
    if not token:
        return None
    s = token.strip()
    if '/' in s:
        xp, yp = s.split('/', 1)
        x = parse_roman_part(xp); y = parse_roman_part(yp)
        if not x or not y:
            return None
        return (y['pc'] + x['pc']) % 12
    p = parse_roman_part(s)
    return p['pc'] if p else None


def ext_suffix(rest, _is_minor):
    """Re-emit extensions (sus/add/7/9/etc.) in the C-major chord-name column."""
    adds = re.findall(r'add(\d+)', rest); rest = re.sub(r'add\d+', '', rest)
    has_maj7 = 'maj7' in rest
    if has_maj7:
        rest = rest.replace('maj7', '')
    main_ext = ''
    m = re.search(r'(?<!\d)(13|11|9|7)(?!\d)', rest)
    if m:
        main_ext = m.group(1); rest = re.sub(r'(?<!\d)(13|11|9|7)(?!\d)', '', rest, count=1)
    sus_str = ''
    susM = re.search(r'sus([24])', rest)
    if susM:
        sus_str = f'sus{susM.group(1)}'; rest = re.sub(r'sus[24]', '', rest)
    is_power = (rest == '5' and not has_maj7 and not main_ext and not sus_str and not adds)
    if is_power:
        rest = ''
    alts = re.findall(r'[#b]\d+', rest)
    result = ''
    if has_maj7: result += 'maj7'
    elif main_ext: result += main_ext
    if sus_str: result += sus_str
    for a in adds: result += f'add{a}'
    for a in alts: result += a
    if is_power: result = '5'
    return result


def chord_to_name(token):
    """Convert a Roman-numeral token to the actual chord name in C major."""
    if not token: return ''
    s = token.strip()
    if '/' in s:
        xp, yp = s.split('/', 1)
        x = parse_roman_part(xp); y = parse_roman_part(yp)
        if not x or not y: return token
        abs_pc = (y['pc'] + x['pc']) % 12
        name = PC_TO_NAME[abs_pc]
        if x['is_minor']: name += 'm'
        name += ext_suffix(x['rest'], x['is_minor'])
        return name
    d = parse_roman_part(s)
    if not d: return token
    name = PC_TO_NAME[d['pc']]
    if d['is_minor']: name += 'm'
    name += ext_suffix(d['rest'], d['is_minor'])
    return name


def style_chord_cell(cell, token):
    """Apply pitch-class-based fill + text color to a chord cell."""
    pc = token_to_pc(token)
    if pc is None:
        return
    pattern, fg, bg, text = PC_STYLE[pc]
    if bg is None:
        cell.fill = PatternFill(patternType=pattern, fgColor=fg)
    else:
        # Diagonal stripes: fg = stripe color, bgColor = field color
        cell.fill = PatternFill(patternType=pattern, fgColor=fg, bgColor=bg)
    cell.font = Font(color=text, bold=True, size=11)
    cell.alignment = Alignment(horizontal='center', vertical='center')


def main():
    with open(CSV_PATH) as f:
        rows = list(csv.DictReader(f))

    wb = Workbook(); ws = wb.active; ws.title = '4-chord progressions'
    headers = ['artist', 'title', 'progression', 'in C / Am',
               'chord 1', 'chord 2', 'chord 3', 'chord 4',
               'sections', 'measures', 'total_repeats', 'key', 'scale']
    ws.append(headers)

    hfill = PatternFill('solid', fgColor='303040')
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.fill = hfill
        cell.font = Font(color='FFFFFF', bold=True, size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for ri, r in enumerate(rows, 2):
        ws.cell(ri, 1, r['artist'])
        ws.cell(ri, 2, r['title'])
        ws.cell(ri, 3, r['progression'])
        ws.cell(ri, 4, '-'.join(chord_to_name(t) for t in r['progression'].split('-')))
        for ci, tok in enumerate(r['progression'].split('-')[:4]):
            cell = ws.cell(ri, 5 + ci, tok)
            style_chord_cell(cell, tok)
        ws.cell(ri, 9, r['sections'])
        ws.cell(ri, 10, r['measures'])
        ws.cell(ri, 11, int(r['total_repeats']) if r['total_repeats'].isdigit() else r['total_repeats'])
        ws.cell(ri, 12, r['key'])
        ws.cell(ri, 13, r['scale'])

    for i, w in enumerate([22, 30, 22, 22, 12, 12, 12, 12, 28, 10, 12, 8, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    wb.save(XLSX_PATH)
    print(f"wrote {XLSX_PATH}  ({len(rows)} rows)")


if __name__ == '__main__':
    main()
