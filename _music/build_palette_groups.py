#!/usr/bin/env python3
"""
Group the guitar-pool songs by how each section OPENS: the palette word of its
first four bars and the shape those chords make.

    themap_venv/python build_palette_groups.py          # whole library -> palette-groups.html + palette-map.html

The unit is the first four bars of every section, starting at its first chord,
leaving out intros and outros (the user: "remove outro and intro from the
discussion"). Two readings of those four bars:

  PALETTE  every chord in them, read in C through the relative major (the nine
           sounds, plus anything else as itself), named by chord_words.json.
           No chord is dropped -- a set with no word is shown as its chords.
  SHAPE    from the chord DURATIONS, not one letter per bar. "C F | G" is
           short-short-long, a back runner, and "C F | G | C F | G" is that
           two-bar back runner twice. The smallest repeating (chord, length)
           cell is scaled to whole slots and repeated to fill four, then named
           from the user's table. A cell that will not fit four slots is shown
           as its lengths rather than forced into a name.
"""
import glob, html, json, math, os, re, sys
from fractions import Fraction
from math import gcd
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
import bars as B
import chord_key as ck
import pao_pools as pp
import build_pool_quiz as Q

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'palette-groups.html')
WORDS = json.load(open(os.path.join(HERE, 'chord_words.json')))
BYSET = {frozenset(v): k for k, v in WORDS.items()}
NINE = {(0, ''): 'C', (2, 'm'): 'Dm', (4, 'm'): 'Em', (5, ''): 'F', (7, ''): 'G',
        (9, 'm'): 'Am', (10, ''): 'A#', (2, ''): 'D', (4, ''): 'E'}
DIATONIC_Q = {0: '', 2: 'm', 4: 'm', 5: '', 7: '', 9: 'm', 11: 'dim', 10: ''}
SHAPES = {'AAAA': 'drone', 'AAAB': 'ender', 'AABA': 'blues', 'AABB': 'shift',
          'ABAA': 'unnamed (ABAA)', 'ABAB': 'vamp', 'ABBA': 'mountain', 'ABBB': 'back runner',
          'AABC': 'front runner', 'ABBC': 'valley', 'ABCC': 'back runner', 'ABAC': 'bounce',
          'ABCB': 'return', 'ABCA': 'mountain', 'ABCD': 'wander'}
ORDER = ['C', 'Dm', 'D', 'Em', 'E', 'F', 'G', 'Am', 'A#']
DEG = {'C': 1, 'Dm': 2, 'D': 2, 'Em': 3, 'E': 3, 'F': 4, 'G': 5, 'Am': 6, 'A#': 7}


def in_c(c, tonic, scale):
    """Hookpad chord -> the chord as it reads in C (relative major). sus, 7, add
    and power chords collapse to their triad -- A and Asus2 are one chord -- but
    nothing is dropped: a chord outside the nine sounds keeps its own name."""
    name, _ = Q.chord_letter(c, tonic, scale)
    root = Q.PC[name[:2]] if name[:2] in Q.PC else Q.PC[name[0]]
    rest = name[len(Q.NAMES[root]):]
    ref = (Q.PC.get(tonic, 0) - Q.PARENT.get(scale, 0)) % 12
    rel = (root - ref) % 12
    if rest.startswith(('dim', '°', 'm7b5')): q = 'dim'
    elif rest.startswith('aug'): q = 'aug'
    elif rest.startswith('m') and not rest.startswith('maj'): q = 'm'
    elif rest.startswith(('sus', '5')): q = DIATONIC_Q.get(rel, '')
    else: q = ''
    return NINE.get((rel, q)) or Q.NAMES[rel] + q


def shape_of(events):
    """[(chord, beats)] -> (name, letters, cell lengths, repeats)"""
    n = len(events)
    cell = events
    for p in range(1, n // 2 + 1):
        if n % p == 0 and all(events[i] == events[i % p] for i in range(n)):
            cell = events[:p]; break
    reps = n // len(cell)
    fr = [Fraction(d).limit_denominator(8) for _, d in cell]
    den = 1
    for f in fr: den = den * f.denominator // gcd(den, f.denominator)
    ints = [int(f * den) for f in fr]
    g = 0
    for x in ints: g = gcd(g, x)
    slots = [x // g for x in ints]
    lengths = '-'.join(f'{float(f):g}' for f in fr)
    k = sum(slots)
    if 4 % k:
        return None, None, lengths, reps
    # A cell shorter than four slots fills them by REPEATING only if it really
    # repeats (F Dm | F Dm ... is a vamp); played once it is STRETCHED, so
    # F | F | G | G stays a shift instead of turning into F G F G
    stretch = 4 // k if reps * k < 4 else 1
    seq, lab = [], {}
    for (ch, _), s in zip(cell, slots):
        lab.setdefault(ch, 'ABCD'[len(lab)] if len(lab) < 4 else '?')
        seq += [lab[ch]] * (s * stretch)
    letters = ''.join(seq) * (4 // len(seq))
    return SHAPES.get(letters), letters, lengths, reps


def openings(path, bpb_override=None):
    d = json.load(open(path))
    bpb = B.beats_per_bar(d)
    keys = d.get('keys') or [{'beat': 1, 'tonic': 'C', 'scale': 'major'}]
    secs = sorted(d.get('sections') or [], key=lambda s: s.get('beat', 1))
    chords = sorted((c for c in d.get('chords') or [] if not c.get('isRest')),
                    key=lambda c: c['beat'])
    for i, s in enumerate(secs):
        nm = str(s.get('name', '')).strip()
        low = nm.lower()
        if low in Q.SKIP or 'intro' in low or 'outro' in low: continue
        lo = s['beat']
        hi = secs[i + 1]['beat'] if i + 1 < len(secs) else (d.get('endBeat') or 10**9)
        inside = [c for c in chords if lo <= c['beat'] < hi]
        if not inside: continue
        st, ss = ck.key_at(keys, lo)
        start = lo + int((inside[0]['beat'] - lo) // bpb) * bpb
        end = min(start + 4 * bpb, hi)
        if end - start < 4 * bpb: continue                  # under four bars
        ev, played = [], []
        for c in inside:
            a, b = max(c['beat'], start), min(c['beat'] + c.get('duration', 0), end)
            if b <= a: continue
            k = in_c(c, st, ss)
            if ev and ev[-1][0] == k: ev[-1][1] += b - a
            else: ev.append([k, b - a])
            name, deg = Q.chord_letter(c, st, ss)
            played.append((c['beat'], b, name, deg))
        # the window must be covered, or the shape is of a fragment
        if sum(x for _, x in ev) < 4 * bpb - 1e-6: continue
        # bars as played, in the file's own key, for reading
        bars = []
        for j in range(4):
            b0, b1 = start + j * bpb, start + (j + 1) * bpb
            bar = []
            for on, off, name, deg in played:
                if on < b1 and off > b0 + 1e-6 and (not bar or bar[-1][0] != name):
                    bar.append([name, deg])
            bars.append(bar)
        cbar = []
        for j in range(4):
            b0, b1 = start + j * bpb, start + (j + 1) * bpb
            bar = []
            for c in inside:
                if c['beat'] < b1 and c['beat'] + c.get('duration', 0) > b0 + 1e-6:
                    k = in_c(c, st, ss)
                    if not bar or bar[-1] != k: bar.append(k)
            cbar.append(bar)
        events = [(k, float(x)) for k, x in ev]
        # Look for a repeating 1- or 2-bar BLOCK before merging across the
        # whole window. Merging first joins the last chord of one pass to the
        # first of the next: Real World's bridge "C G | F C" x2 became
        # C G F C(long) G F C and lost both the repeat and its mountain.
        def span(a, b):
            out = []
            for c in inside:
                x0, x1 = max(c['beat'], a), min(c['beat'] + c.get('duration', 0), b)
                if x1 <= x0: continue
                k = in_c(c, st, ss)
                if out and out[-1][0] == k: out[-1][1] += x1 - x0
                else: out.append([k, x1 - x0])
            return [(k, float(x)) for k, x in out]
        block_reps = 1
        for p in (1, 2):
            blocks = [span(start + i * p * bpb, start + (i + 1) * p * bpb) for i in range(4 // p)]
            if all(b == blocks[0] for b in blocks) and len(blocks[0]) > 1:
                events, block_reps = blocks[0], 4 // p
                break
        # A PUSHED change (landing on an "and", Dancing in the Dark's 3.5 + 4.5)
        # is rhythm texture, not shape: snap each change to the beat it
        # anticipates for naming -- unless that would swallow a chord
        snapped, t, pushed = [], 0.0, False
        for k, x in events:
            on, off = t, t + x
            t = off
            s_on, s_off = math.floor(on + 0.5), math.floor(off + 0.5)
            pushed |= (s_on != on)
            snapped.append((k, s_off - s_on))
        if all(x > 0 for _, x in snapped) and pushed:
            events, push = snapped, True
        else:
            push = False
        # judge the shape on every pass of the block, kept as separate entries,
        # so a repeating "F Dm" is still a vamp and not one stretched pass
        name, letters, lengths, reps = shape_of(events * block_reps)
        pal = sorted({k for k, _ in events}, key=lambda x: (ORDER.index(x) if x in ORDER else 99, x))
        yield dict(section=nm, key=f'{st} {ss}', bars=bars, cbars=cbar, palette=pal,
                   word=BYSET.get(frozenset(pal)), shape=name, letters=letters,
                   lengths=lengths, reps=reps, pushed=push, cellbars=round(4 / reps, 2))


GROUPS = ['G50-250', 'G300-650', 'G700-1300', 'Beatles', 'other']
SCRATCH = re.compile(r'^(music_|perms_|mine|chord-riffs|\d+-\d+-\d+)', re.I)


def group(lst, stem):
    if lst:
        n = int(lst[1:])
        return 'G50-250' if n <= 250 else 'G300-650' if n <= 650 else 'G700-1300'
    return 'Beatles' if stem.startswith('beatles_') else 'other'


def listed_songs(wb, files):
    """{file: (list, title, artist)} for every song on a Guitar sheet. Shared
    with build_tempo_map."""
    # guitar lists, smallest first: a song on several keeps the first
    sheets = sorted((n for n in wb.sheetnames if re.fullmatch(r'Guitar\s*\d+', n)),
                    key=lambda n: int(re.sub(r'\D', '', n)))
    listed = {}
    for sheet in sheets:
        lst = 'G' + re.sub(r'\D', '', sheet)
        sr = [r for r in wb[sheet].iter_rows(values_only=True) if any(r)]
        head = [str(h or '').strip().lower() for h in sr[0]]
        if {'artist', 'band'} & set(head) and {'title', 'song'} & set(head):
            ia = head.index('artist') if 'artist' in head else head.index('band')
            it = head.index('title') if 'title' in head else head.index('song')
            pairs, check = [(r[ia], r[it]) for r in sr[1:]], False
        else:
            # G300+ have no header and swap the columns from sheet to sheet:
            # take the first two text cells, try both ways round, and keep the
            # one whose matched file has that artist
            pairs, check = [], True
            for r in sr:
                t = [c for c in r if isinstance(c, str) and c.strip()][:2]
                if len(t) == 2: pairs += [(t[0], t[1]), (t[1], t[0])]
        for artist, title in pairs:
            if not isinstance(artist, str) or not isinstance(title, str): continue   # blank / stray rows (G50 has some)
            path = Q.resolve(artist, title, files)
            if not path or path in listed: continue
            fa = pp.norm(os.path.basename(path).split('_', 1)[0])
            if check and fa and pp.norm(artist)[:4] not in fa and fa[:4] not in pp.norm(artist): continue
            listed[path] = (lst, str(title), str(artist))
    return listed


def library_files(files, listed):
    """the songs the whole-library pages cover: not the user's own, the
    workbenches, or -variants (unless a guitar list names the variant)"""
    out = []
    for path in sorted(files):
        stem = os.path.splitext(os.path.basename(path))[0]
        if SCRATCH.match(stem) or re.search(r'-(midi|pasted)$', stem): continue
        if path not in listed and pp.split_variant(stem)[1]: continue
        out.append(path)
    return out


def main():
    """Every song in the Hookpad library (not the user's own, the workbenches or
    -variants). A song on a guitar list keeps that list's title/artist and its
    list as a tag; the rest are named from the file."""
    wb = openpyxl.load_workbook(Q.TRACKER, read_only=True)
    library = json.load(open(Q.HOOKPAD_LIST))
    live = {s.get('song', '').lower() for s in library}
    files = [f for f in glob.glob(os.path.join(pp.EXPORT, '*.json'))
             if os.path.splitext(os.path.basename(f))[0].lower() in live]
    listed = listed_songs(wb, files)
    rows = []
    for path in library_files(files, listed):
        stem = os.path.splitext(os.path.basename(path))[0]
        lst, title, artist = listed.get(path) or (None, *(stem.split('_', 1)[::-1] if '_' in stem else (stem, '')))
        if title == title.lower(): title = Q.title_case(title)
        if artist == artist.lower(): artist = Q.title_case(artist)
        hp = Q.hookpad_url(path, library)
        seen = {}
        try: ops = list(openings(path))
        except Exception: continue
        for o in ops:
            k = json.dumps(o['cbars'])
            if k in seen:                       # same opening, another section
                if o['section'] not in seen[k]['section']:
                    seen[k]['section'] += ' = ' + o['section']
                continue
            o.update(pool=group(lst, stem), list=lst or '', title=title, artist=artist, hookpad=hp)
            seen[k] = o
            rows.append(o)
    for name, out in (('palette_groups_template.html', OUT), ('palette_map_template.html', os.path.join(HERE, 'palette-map.html'))):
        page = open(os.path.join(HERE, name)).read()
        open(out, 'w').write(page.replace('__ROWS__', json.dumps(rows, ensure_ascii=False))
                                 .replace('__WORDS__', json.dumps(WORDS)).replace('__POOLS__', json.dumps(GROUPS)))
    named = sum(1 for r in rows if r['shape'])
    print(f'{OUT}\n  {len(rows)} openings from {len({(r["artist"], r["title"]) for r in rows})} songs'
          f'\n  palette word: {sum(1 for r in rows if r["word"])}   shape name: {named}'
          f'   lengths only: {len(rows) - named}')
    for g in GROUPS: print(f'  {g:10} {len({(r["artist"], r["title"]) for r in rows if r["pool"] == g})} songs')


if __name__ == '__main__':
    main()
