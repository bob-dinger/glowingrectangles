#!/usr/bin/env python3
"""
Every front runner (long-short-short, AABC) and back runner (short-short-long,
ABCC) in the library, and the chords they use -> ~/Desktop/runners.xlsx

    themap_venv/python runners_sheet.py

Shapes come from build_palette_groups (chord durations, scale-free: "F Am | G"
is a back runner). Every four-bar block of every section is read, not just the
opening, so a runner in bars 5-8 counts too. Intros/outros, the user's own songs
(mine/perms/music_) and variants are left out. Chords are read in C through the
relative major, so Too Little Too Late's F Am | G sits with every other F Am G.

Also writes _music/runners-map.html (runners_map_template.html): one circle per
chord order, sized by songs, joined to the same loop heard from another bar.

Sheets: Combos (one row per chord order, with songs), Coverage (every order of
three of the six diatonic chords, found or not), Sections (every instance).
"""
import collections, glob, itertools, json, os, re, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import build_palette_groups as P
import bars as B
from hashids import Hashids

D = os.path.expanduser('~/Desktop/music/hookpad_songs_full')
OUT = os.path.expanduser('~/Desktop/runners.xlsx')
SCRATCH = re.compile(r'^(music_|perms_|mine|chord-riffs|\d+-\d+-\d+)', re.I)
VARIANT = re.compile(r'-(\d{2,3}|hooktab|simple|double|C|right|wrong|mixolydian|filled|3way)$')
HASH = Hashids(salt='XI0Y4UFrK6EPLnarrI4y', min_length=11, alphabet='ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz-_')
ROMAN = {'C': 'I', 'Dm': 'ii', 'D': 'II', 'Em': 'iii', 'E': 'III', 'F': 'IV', 'G': 'V', 'Am': 'vi', 'A#': 'bVII'}
SIX = ['C', 'Dm', 'Em', 'F', 'G', 'Am']
WANT = {'AABC': 'front runner', 'ABCC': 'back runner'}


def blocks(path):
    """the song with every section cut into four-bar pieces, as openings() reads them"""
    d = json.load(open(path))
    bpb = B.beats_per_bar(d)
    secs = sorted(d.get('sections') or [], key=lambda s: s.get('beat', 1))
    end = d.get('endBeat') or 0
    chords = [c for c in d.get('chords') or [] if not c.get('isRest')]
    out = []
    for i, s in enumerate(secs):
        lo, hi = s['beat'], secs[i + 1]['beat'] if i + 1 < len(secs) else end
        inside = sorted(c['beat'] for c in chords if lo <= c['beat'] < hi)
        if not inside: continue
        a = lo + int((inside[0] - lo) // bpb) * bpb
        nm = str(s.get('name', '')).strip()
        if nm.lower() in P.Q.SKIP: continue
        while a + 4 * bpb <= hi:   # the bar number rides on the name, openings() passes it through
            out.append({'name': f'{nm} @{int((a - 1) // bpb) + 1}', 'beat': a}); a += 4 * bpb
    d['sections'] = out
    return d, bpb


def order(o):
    """the chords in role order A, B, C"""
    seen = []
    for bar in o['cbars']:
        for c in bar:
            if c not in seen: seen.append(c)
    return seen


def main():
    ids = {}
    for s in json.load(open(os.path.expanduser('~/Desktop/music/.hookpad_song_list.json'))):
        ids.setdefault(s['song'].lower(), s['ID'])
    rows = []
    tmp = os.path.join(tempfile.mkdtemp(), 's.json')
    for f in sorted(glob.glob(D + '/*.json')):
        song = os.path.basename(f)[:-5]
        if SCRATCH.match(song) or VARIANT.search(song): continue
        try:
            d, bpb = blocks(f)
            json.dump(d, open(tmp, 'w'))
            ops = list(P.openings(tmp))
        except Exception:
            continue
        for o in ops:
            if o['letters'] not in WANT: continue
            ch = order(o)
            if len(ch) != 3: continue
            sec, bar = o['section'].rsplit(' @', 1)
            rows.append(dict(song=song, section=sec, bar=int(bar), shape=WANT[o['letters']], chords=ch,
                             played=' | '.join(' '.join(c for c, _ in bar) for bar in o['bars']),
                             inc=' | '.join(' '.join(bar) for bar in o['cbars']), cbars=o['cbars'], key=o['key'],
                             cell=o['cellbars'], lengths=o['lengths'], pushed=o['pushed'],
                             url=f'https://hookpad.hooktheory.com/?idOfUserSong={HASH.encode(ids[song.lower()])}' if song.lower() in ids else ''))
    # one instance per (song, shape, chord order): a repeated verse is one entry
    seen, uniq = set(), []
    for r in rows:
        k = (r['song'], r['shape'], tuple(r['chords']))
        if k not in seen: seen.add(k); uniq.append(r)
    combos = collections.defaultdict(list)
    for r in uniq: combos[(r['shape'], tuple(r['chords']))].append(r)

    wb = openpyxl.Workbook()
    bold, hdr = Font(bold=True), PatternFill('solid', fgColor='DDDDDD')
    def sheet(ws, head, widths):
        ws.append(head)
        for c in ws[1]: c.font, c.fill = bold, hdr
        for i, w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = 'A2'
    ws = wb.active; ws.title = 'Combos'
    sheet(ws, ['shape', 'chords (in C)', 'roman', 'how it sounds', 'songs', 'all diatonic?', 'songs (first 12)'], [13, 16, 16, 18, 7, 12, 120])
    for (shape, ch), rs in sorted(combos.items(), key=lambda kv: (kv[0][0], -len({r['song'] for r in kv[1]}))):
        a, b, c = ch
        sound = f'{a} | {b} {c}' if shape == 'front runner' else f'{a} {b} | {c}'
        songs = sorted({r['song'] for r in rs})
        ws.append([shape, ' '.join(ch), ' '.join(ROMAN.get(x, x) for x in ch), sound, len(songs),
                   'yes' if all(x in SIX for x in ch) else 'no', ', '.join(songs[:12]) + (' ...' if len(songs) > 12 else '')])
    ws = wb.create_sheet('Coverage')
    sheet(ws, ['shape', 'chords (in C)', 'roman', 'songs', 'found?'], [13, 16, 16, 7, 8])
    miss = PatternFill('solid', fgColor='F6D5D5')
    for shape in ('front runner', 'back runner'):
        allp = sorted(itertools.permutations(SIX, 3), key=lambda p: -len({r['song'] for r in combos.get((shape, p), [])}))
        for p in allp:
            n = len({r['song'] for r in combos.get((shape, p), [])})
            ws.append([shape, ' '.join(p), ' '.join(ROMAN[x] for x in p), n, 'yes' if n else 'no'])
            if not n:
                for c in ws[ws.max_row]: c.fill = miss
    ws = wb.create_sheet('Sections')
    sheet(ws, ['song', 'section', 'bar', 'shape', 'chords (in C)', 'as played', 'in C', 'key', 'cell (bars)', 'pushed', 'hookpad'], [42, 14, 5, 13, 14, 30, 30, 12, 9, 7, 60])
    for r in sorted(uniq, key=lambda r: (r['shape'], ' '.join(r['chords']), r['song'])):
        ws.append([r['song'], r['section'], r['bar'], r['shape'], ' '.join(r['chords']), r['played'], r['inc'], r['key'], r['cell'], 'yes' if r['pushed'] else '', r['url']])
    wb.save(OUT)
    # the same instances as a map: one circle per chord order -> runners-map.html
    page = []
    for r in uniq:
        artist, title = r['song'].split('_', 1) if '_' in r['song'] else ('', r['song'])
        page.append(dict(song=r['song'], title=P.Q.title_case(title), artist=P.Q.title_case(artist), section=r['section'],
                         bar=r['bar'], shape=r['shape'], chords=r['chords'], cbars=r['cbars'], key=r['key'],
                         cell=r['cell'], pushed=r['pushed'], url=r['url']))
    here = os.path.dirname(os.path.abspath(__file__))
    tpl = open(os.path.join(here, 'runners_map_template.html')).read()
    open(os.path.join(here, 'runners-map.html'), 'w').write(tpl.replace('__ROWS__', json.dumps(page, ensure_ascii=False)))
    for shape in ('front runner', 'back runner'):
        cs = [(k[1], len({r['song'] for r in v})) for k, v in combos.items() if k[0] == shape]
        dia = [c for c in cs if all(x in SIX for x in c[0])]
        print(f'{shape}: {sum(n for _, n in cs)} song-instances, {len(cs)} chord orders; diatonic {len(dia)}/120 found')
        for c, n in sorted(cs, key=lambda x: -x[1])[:12]: print(f'   {" ".join(c):14} {n}')
    print('jojo:', [(r['section'], r['shape'], r['inc']) for r in uniq if r['song'].startswith('jojo_too little')])
    print(OUT)


if __name__ == '__main__':
    main()
