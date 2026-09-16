"""Generate ~/Desktop/section_curation.xlsx — one row per section across
Beatles-Study + Guitar50 + Guitar100. Seed columns auto-filled from
Hookpad data; template / notes left blank for hand curation.
"""
import os, re, json
from collections import Counter
from dotenv import load_dotenv
load_dotenv('/Users/robert/Desktop/themap/themap_claude/.env')
from supabase import create_client
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PAGES = ['Beatles-Study.html', 'Guitar50.html', 'Guitar100.html']
OUT = os.path.expanduser('~/Desktop/section_curation.xlsx')

QUAL = {'major':['M','m','m','M','M','m','d'],'minor':['m','d','M','m','m','M','M']}
ROMAN_UP = ['I','II','III','IV','V','VI','VII']
ROMAN_LO = ['i','ii','iii','iv','v','vi','vii']


def chord_label(c, scale='major'):
    root = str(c.get('root',''))
    acc, rs = '', root
    while rs and rs[0] in 'b#': acc += rs[0]; rs = rs[1:]
    if not rs.isdigit(): return '?'
    deg = int(rs)
    if c.get('applied'): q = 'M'
    elif c.get('type') in ('m','min'): q = 'm'
    else: q = QUAL.get(scale, QUAL['major'])[deg-1]
    r = (ROMAN_UP if q in ('M','a') else ROMAN_LO)[deg-1]
    lbl = acc + r
    if c.get('type') in (7,'7'): lbl += '7'
    if q == 'd': lbl += '°'
    if c.get('applied'): lbl += f'/{ROMAN_UP[c["applied"]-1]}'
    return lbl


def slugs_from_page(fn):
    html = open(os.path.join(HERE, fn)).read()
    m = re.search(r'const SONGS\s*=\s*(\[.*?\]);', html, flags=re.DOTALL)
    return [(s['slug'], s.get('title',''), s.get('artist',''))
            for s in json.loads(m.group(1)) if s.get('slug')]


def chords_per_measure(start, n_bars, sec_chords, bpb, scale):
    measures = [[] for _ in range(n_bars)]
    for c in sec_chords:
        idx = int((c.get('beat', 0) - start) / bpb)
        if 0 <= idx < n_bars:
            measures[idx].append(chord_label(c, scale))
    out, last = [], None
    for m in measures:
        if m:
            out.append(' '.join(m)); last = m[-1]
        else:
            out.append(last or '—')
    return out


def phrase_letter_guess(pmc):
    """Greedy: split into 2-bar chunks, letter by first-occurrence equality."""
    n = len(pmc)
    if n < 2 or n % 2: return ''
    chunks = [tuple(pmc[i:i+2]) for i in range(0, n, 2)]
    letters, seen = [], {}
    for ch in chunks:
        if ch not in seen:
            seen[ch] = chr(ord('A') + len(seen))
        letters.append(seen[ch])
    return ''.join(letters)


def section_breakdown(d):
    sections = d.get('sections') or []
    bpb = ((d.get('meters') or [{}])[0].get('numBeats')) or 4
    scale = (d.get('keys') or [{}])[0].get('scale') or 'major'
    end_beat = d.get('endBeat') or 0
    chords = d.get('chords') or []
    sorted_secs = sorted(sections, key=lambda s: s.get('beat', 0))
    out = []
    for i, s in enumerate(sorted_secs):
        name = (s.get('name') or '').strip()
        if not name or name.lower() == 'section':
            continue
        start = s.get('beat', 0)
        end = sorted_secs[i+1]['beat'] if i+1 < len(sorted_secs) else end_beat
        bars = round((end - start) / bpb)
        if bars < 1:
            continue
        sec_ch = [c for c in chords if start <= c.get('beat', 0) < end and not c.get('isRest')]
        pmc = chords_per_measure(start, bars, sec_ch, bpb, scale)
        out.append({
            'name': name, 'bars': bars,
            'pmc': pmc,
            'letters': phrase_letter_guess(pmc),
        })
    return out


def build():
    sb = create_client(os.environ['SUPABASE_URL'], os.environ['SUPABASE_SERVICE_ROLE_KEY'])
    pool = {}
    for fn in PAGES:
        page_name = fn.replace('.html','')
        for slug, title, artist in slugs_from_page(fn):
            pool.setdefault(slug, {'title': title, 'artist': artist, 'pages': []})
            pool[slug]['pages'].append(page_name)
    slugs = list(pool.keys())
    print(f"pool: {len(slugs)} unique slugs across {PAGES}")

    rows = []
    for i in range(0, len(slugs), 100):
        chunk = slugs[i:i+100]
        data = (sb.schema('parcels').table('songs')
                .select('slug,hookpad_json').in_('slug', chunk).execute().data)
        by_slug = {r['slug']: r for r in data}
        for slug in chunk:
            meta = pool[slug]
            r = by_slug.get(slug)
            if not r or not r.get('hookpad_json'):
                continue
            for sec in section_breakdown(r['hookpad_json']):
                rows.append({
                    'pool': ' · '.join(meta['pages']),
                    'artist': meta['artist'],
                    'title': meta['title'],
                    'slug': slug,
                    'section': sec['name'],
                    'bars': sec['bars'],
                    'letters_guess': sec['letters'],
                    'template_guess': f"{sec['bars']}-{len(sec['pmc'])}-{sec['letters']}" if sec['letters'] else f"{sec['bars']}",
                    'chords_per_measure': ' | '.join(sec['pmc']),
                    'template': '',
                    'notes': '',
                })
    print(f"sections: {len(rows)}")

    wb = Workbook()
    ws = wb.active
    ws.title = 'sections'
    headers = ['pool','artist','title','section','bars','letters_guess',
               'template_guess','template','notes','chords_per_measure','slug']
    ws.append(headers)
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='6366f1')
    for col_i, h in enumerate(headers, 1):
        c = ws.cell(1, col_i)
        c.font = head_font; c.fill = head_fill
        c.alignment = Alignment(horizontal='left', vertical='center')
    for row in rows:
        ws.append([row.get(h, '') for h in headers])
    widths = {'pool': 18, 'artist': 20, 'title': 28, 'section': 14, 'bars': 6,
              'letters_guess': 12, 'template_guess': 14, 'template': 14,
              'notes': 24, 'chords_per_measure': 60, 'slug': 38}
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUT)
    print(f"wrote {OUT}")
    # Quick summary of template_guess distribution
    counter = Counter(r['template_guess'] for r in rows)
    print(f"\nTop 20 auto-guessed templates:")
    for tmpl, n in counter.most_common(20):
        print(f"  {n:>4}  {tmpl}")


if __name__ == '__main__':
    build()
