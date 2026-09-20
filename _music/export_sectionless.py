#!/usr/bin/env python3
"""The live Hookpad songs with no section markers, as a worklist.

Sorted so the ones worth doing first come first: a song that already has chords
and a melody but no sections is otherwise complete, and sectioning it is the one
thing standing between it and every structural analysis in this folder.
"""
import json, glob, os, re
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

D = os.path.expanduser('~/Desktop/music/hookpad_songs_full')

SCRATCH = re.compile(r'^(music_|perms_|mine_\d|\d+-\d+-\d+)', re.I)
def is_scratch(n): return bool(SCRATCH.match(n))

live = {s['song'].replace('/', '_').lower()
        for s in json.load(open(os.path.expanduser('~/Desktop/music/.hookpad_song_list.json')))}

rows = []
for f in glob.glob(f"{D}/*.json"):
    n = os.path.basename(f)[:-5]
    if n.lower() not in live or is_scratch(n): continue
    try: d = json.load(open(f))
    except Exception: continue
    if d.get('sections'): continue
    ch = d.get('chords') or []
    notes = d.get('notes') or []
    nb = (d.get('meters') or [{}])[0].get('numBeats', 4) or 4
    ends = [c['beat'] + c['duration'] for c in ch] + [x['beat'] + x['duration'] for x in notes]
    bars = round(max(ends) / nb) if ends else 0
    k = (d.get('keys') or [{}])[0]
    k = k if isinstance(k, dict) else {}
    a, _, t = n.partition('_')
    rows.append([a.title(), t.title(), n, len(ch), len(notes), bars,
                 f"{k.get('tonic') or ''} {k.get('scale') or ''}".strip(),
                 (d.get('tempos') or [{}])[0].get('bpm') and
                 round((d.get('tempos') or [{}])[0]['bpm']),
                 f"{nb}/4",
                 "ready — chords + melody" if ch and notes else
                 ("chords only" if ch else ("melody only" if notes else "empty"))])

# most complete first, then longest
order = {"ready — chords + melody": 0, "chords only": 1, "melody only": 2, "empty": 3}
rows.sort(key=lambda r: (order[r[9]], -r[5], r[0], r[1]))

wb = Workbook(); ws = wb.active; ws.title = "No sections"
cols = [("Artist",24),("Title",40),("Hookpad file",48),("Chords",8),("Melody notes",13),
        ("Bars",7),("Key",13),("BPM",7),("Meter",7),("State",24)]
ws.append([c for c, _ in cols])
for i,(_,w) in enumerate(cols,1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]: c.font = Font(bold=True)
ws.freeze_panes = "A2"
for r in rows: ws.append(r)
out = os.path.expanduser('~/Desktop/hookpad_no_sections.xlsx')
wb.save(out)

import collections
c = collections.Counter(r[9] for r in rows)
print(f"  {len(rows):,} live songs with no sections")
for k, v in sorted(c.items(), key=lambda kv: order[kv[0]]):
    print(f"    {k:<24}{v:>5}")
print(f"  {out}")
