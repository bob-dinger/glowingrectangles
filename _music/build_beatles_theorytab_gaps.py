"""Cross-reference HookTheory TheoryTab's Beatles catalog against the user's
Hookpad library, grouped by Pollack album/session. Writes an actionable xlsx of
songs available in TheoryTab that aren't yet in Hookpad — each with its TheoryTab
URL (use the site's "Open in Hookpad" to import).

    ~/Desktop/beatles_theorytab_to_add.xlsx
"""
import os, re, json, requests
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def norm(s): return re.sub(r'[^a-z0-9]', '', (s or '').lower())
UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/145.0.0.0 Safari/537.36'
POL = os.path.expanduser('~/Desktop/music/pollack_beatles_notes')

ALBUM_ORDER = ["Please Please Me", "With The Beatles", "A Hard Day's Night",
    "Beatles For Sale", "Help!", "Rubber Soul", "Revolver",
    "Sgt. Pepper's Lonely Hearts Club Band", "Magical Mystery Tour",
    "White Album", "Yellow Submarine", "Abbey Road", "Let It Be", "Past Masters",
    "(cover / not in Pollack)"]

def theorytab_songs():
    r = requests.get('https://www.hooktheory.com/theorytab/artist/the-beatles',
                     headers={'User-Agent': UA}, timeout=15)
    slugs = sorted(set(re.findall(r'/theorytab/view/the-beatles/([a-z0-9-]+)', r.text)))
    return {norm(s.replace('-', ' ')): s for s in slugs}     # normkey -> slug

def hookpad_have():
    cat = json.load(open(os.path.expanduser('~/Desktop/music/.hookpad_song_list.json')))
    tp = lambda n: n.split('_', 1)[1] if '_' in n else n
    return {norm(tp(s['song'])) for s in cat if s['song'].lower().startswith('beatles')}

def pollack_albums():
    idx = json.load(open(POL + '/_index.json'))
    album = {}
    for key, meta in idx.items():
        html = open(os.path.join(POL, key + '.html'), encoding='utf-8', errors='replace').read()
        m = re.search(r'CD:\s*"([^"]+)"', html)
        if m:
            a = m.group(1)
            a = 'Past Masters' if 'Past Master' in a else a
            album[norm(meta['title'])] = a
    return album

def main():
    tt, have, album = theorytab_songs(), hookpad_have(), pollack_albums()
    rows = []
    for nk, slug in tt.items():
        if nk in have: continue
        if any(nk in h or h in nk for h in have if len(h) > 5 and len(nk) > 5): continue
        title = slug.replace('-', ' ').strip()
        rows.append((album.get(nk, '(cover / not in Pollack)'), title, slug))

    order = {a: i for i, a in enumerate(ALBUM_ORDER)}
    rows.sort(key=lambda r: (order.get(r[0], 99), r[1]))

    wb = Workbook(); ws = wb.active; ws.title = 'TheoryTab to add'
    ws.append(['Album / Session', 'Song', 'TheoryTab link (Open in Hookpad)', 'Added?'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='7030A0')
        cell.alignment = Alignment(horizontal='center')
    alt = PatternFill('solid', fgColor='EDE7F6')
    last_album, shade = None, False
    for a, title, slug in rows:
        if a != last_album: shade = not shade; last_album = a
        ws.append([a, title, f'https://www.hooktheory.com/theorytab/view/the-beatles/{slug}', ''])
        if shade:
            for col in range(1, 5): ws.cell(ws.max_row, col).fill = alt
    for i, w in enumerate([38, 34, 60, 8], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    out = os.path.expanduser('~/Desktop/beatles_theorytab_to_add.xlsx')
    wb.save(out)
    print(f'wrote {out}  ({len(rows)} songs to add)')
    bya = defaultdict(int)
    for a, _, _ in rows: bya[a] += 1
    for a in ALBUM_ORDER:
        if bya.get(a): print(f'  {a}: {bya[a]}')

if __name__ == '__main__':
    main()
