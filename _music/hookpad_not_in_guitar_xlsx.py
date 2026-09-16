"""List Hookpad songs NOT in any Guitar sheet (50-1300) of songs_to_learn.xlsx, excluding Beatles.

Reads guitar lists directly from the xlsx (so it stays current as the user adds sheets).

Sheet formats vary:
  - early sheets (Guitar50..Guitar650): row = (title, artist, key, ...)
  - later sheets (Guitar700+):          row = (artist, title, "artist_title" hookpad_name)
Detected per-row: a 3rd cell containing '_' => later format.

Inputs:
    ~/Desktop/music/songs_to_learn.xlsx
    ~/Desktop/hookpad_song_names.csv   (from dump_hookpad_names.py)
Output:
    ~/Desktop/hookpad_not_in_guitar.csv
"""
import csv, os, re
from collections import defaultdict
import openpyxl

XLSX = os.path.expanduser('~/Desktop/music/songs_to_learn.xlsx')
HP = os.path.expanduser('~/Desktop/hookpad_song_names.csv')
OUT = os.path.expanduser('~/Desktop/hookpad_not_in_guitar.csv')


def norm(s):
    s = str(s or '').lower()
    s = re.sub(r"['']", '', s)
    s = re.sub(r'[^a-z0-9]+', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


_ARTIST_ALIASES = {
    'creedence clearwater revival': 'ccr',
    'red hot chili peppers': 'rhcp',
    'lynyrd skynyrd': 'lynryd skynyrd',
    'the band': 'band',
    'guns n roses': 'gnr',
}

def norm_artist(s):
    n = norm(s)
    if n.startswith('the '): n = n[4:]
    return _ARTIST_ALIASES.get(n, n)


def artist_overlap(a1, a2):
    if not a1 or not a2: return False
    if a1 == a2 or a1 in a2 or a2 in a1: return True
    w1 = {w for w in a1.split() if len(w) >= 3}
    w2 = {w for w in a2.split() if len(w) >= 3}
    return bool(w1 & w2)


def load_guitar():
    """Return (by_title index, set of (na,nt), count of sheets)."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    by_title = defaultdict(list)
    total = set()
    sheets = [s for s in wb.sheetnames if s.lower().replace(' ', '').startswith('guitar')]
    for name in sheets:
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            if not row or all(c is None for c in row): continue
            c0, c1 = row[0], row[1] if len(row) > 1 else None
            c2 = row[2] if len(row) > 2 else None
            # later format: 3rd cell is "artist_title"
            if isinstance(c2, str) and '_' in c2 and len(c2.split('_', 1)) == 2:
                art, tit = c2.split('_', 1)
            else:
                tit, art = c0, c1
            t, a = norm(tit), norm_artist(art)
            if t:
                by_title[t].append((a, f'{art} - {tit}', name))
                total.add((a, t))
    return by_title, total, len(sheets)


def in_guitar(nart, ntit, by_title):
    for a, _, _ in by_title.get(ntit, []):
        if artist_overlap(nart, a): return True
    tc = ntit.replace(' ', '')
    if len(tc) >= 5:
        for gt, lst in by_title.items():
            gc = gt.replace(' ', '')
            if tc == gc or (abs(len(tc) - len(gc)) <= 1 and (tc in gc or gc in tc)):
                for a, _, _ in lst:
                    if artist_overlap(nart, a): return True
    tw = {w for w in ntit.split() if len(w) >= 3}
    if tw:
        need = max(1, int(0.7 * len(tw)))
        for gt, lst in by_title.items():
            gw = {w for w in gt.split() if len(w) >= 3}
            if gw and len(tw & gw) >= need:
                for a, _, _ in lst:
                    if artist_overlap(nart, a): return True
    return False


# User's own scratch/sketch Hookpad files (not real songs to learn)
# normalized artist forms (norm() turns '-' into space)
SCRATCH = {'mine', 'song', 'music', 'riffs', 'chord riffs', 'ship to wreck',
           'shiptowreck', 'verse', 'chorus'}

def is_scratch(na, name):
    if not na: return True                          # blank artist (mine1, mrtambourine, ...)
    if re.fullmatch(r'mine\d*[a-z]?', na): return True  # mine, mine1..mine22, mine18a
    return na in SCRATCH


def main():
    by_title, total, n_sheets = load_guitar()
    print(f'{len(total)} unique songs across {n_sheets} guitar sheets')

    seen = set(); hp = []; n_beatles = 0; n_scratch = 0
    for r in csv.DictReader(open(HP)):
        name = (r.get('name') or '').strip()
        if not name: continue
        parts = name.split('_', 1)
        art, tit = (parts[0], parts[1]) if len(parts) == 2 else ('', name)
        na, nt = norm_artist(art), norm(tit)
        if 'beatles' in na or 'beatles' in norm(name):
            n_beatles += 1; continue
        if is_scratch(na, name):
            n_scratch += 1; continue
        key = (na, nt)
        if key in seen: continue
        seen.add(key)
        hp.append({'artist': art, 'title': tit, 'na': na, 'nt': nt, 'name': name})

    print(f'{len(hp)} unique non-Beatles Hookpad songs ({n_beatles} Beatles, {n_scratch} scratch/sketch rows skipped)')

    missing = [s for s in hp if not in_guitar(s['na'], s['nt'], by_title)]
    missing.sort(key=lambda s: (s['na'], s['nt']))
    print(f'\n{len(missing)} Hookpad songs NOT in any guitar sheet:\n')

    with open(OUT, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['artist', 'title', 'hookpad_name'])
        for s in missing:
            w.writerow([s['artist'], s['title'], s['name']])
            print(f'  {s["artist"]} — {s["title"]}')

    print(f'\nfull list ({len(missing)} songs) -> {OUT}')


if __name__ == '__main__': main()
