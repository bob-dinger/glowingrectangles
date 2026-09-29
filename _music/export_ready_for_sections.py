#!/usr/bin/env python3
"""Just the songs that need ONLY section markers.

Chords in, melody in, no sections. Each is one pass of marker-dropping away
from entering the phrase census, the window census and the block system.

Deliberately a separate file from hookpad_no_sections.xlsx, which the user
edits by hand -- this one is safe to regenerate, that one is not.
"""
import json, glob, os, re
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

D = os.path.expanduser('~/Desktop/music/hookpad_songs_full')
SCRATCH = re.compile(r'^(music_|perms_|mine|\d+-\d+-\d+)', re.I)
live = {s['song'].replace('/', '_').lower()
        for s in json.load(open(os.path.expanduser('~/Desktop/music/.hookpad_song_list.json')))}

rows = []
for f in glob.glob(f"{D}/*.json"):
    n = os.path.basename(f)[:-5]
    if n.lower() not in live or SCRATCH.match(n): continue
    try: d = json.load(open(f))
    except Exception: continue
    if d.get('sections'): continue
    ch, notes = d.get('chords') or [], d.get('notes') or []
    if not (ch and notes): continue                 # the "ready" test
    nb = (d.get('meters') or [{}])[0].get('numBeats', 4) or 4
    ends = [c['beat']+c['duration'] for c in ch] + [x['beat']+x['duration'] for x in notes]
    k = (d.get('keys') or [{}])[0]; k = k if isinstance(k, dict) else {}
    bpm = (d.get('tempos') or [{}])[0].get('bpm')
    a, _, t = n.partition('_')
    rows.append([a.title(), t.title(), n, round(max(ends)/nb), len(ch), len(notes),
                 f"{k.get('tonic') or ''} {k.get('scale') or ''}".strip(),
                 bpm and round(bpm), f"{nb}/4", ""])

rows.sort(key=lambda r: (-r[3], r[0], r[1]))        # longest first = most song there

wb = Workbook(); ws = wb.active; ws.title = "Ready for sections"
cols = [("Artist",24),("Title",40),("Hookpad file",46),("Bars",7),("Chords",8),
        ("Melody notes",13),("Key",13),("BPM",7),("Meter",7),("Done",8)]
ws.append([c for c,_ in cols])
for i,(_,w) in enumerate(cols,1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]: c.font = Font(bold=True)
ws.freeze_panes = "A2"
for r in rows: ws.append(r)
out = os.path.expanduser('~/Desktop/hookpad_ready_for_sections.xlsx')
wb.save(out)
print(f"  {len(rows)} songs need only section markers")
print(f"  {out}")
