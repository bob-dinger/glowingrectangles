#!/usr/bin/env python3
"""
Guitar pool -> spreadsheet of song / band / part / progression / PAO variants.

Each row is one section of one song. The PAO columns are independent random
scenes for the same progression, so you can pick whichever suits the song —
the point being that every part gets its own picture.

    python3 pao_sheet.py                       # G50, 3 variants
    python3 pao_sheet.py --pools G50,G100 --variants 5
"""
import argparse, json, os, random, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pao_pools as pp
import pao_options as po

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill('solid', fgColor='1F2430')
HF   = Font(color='FFFFFF', bold=True, size=10)


def scene_for(seq, rng):
    """One full scene (possibly several sentences for long progressions)."""
    out = []
    for i in range(0, len(seq), 4):
        grp = seq[i:i+4]
        picks = []
        for slot, ch in enumerate(grp):
            # applied dominants normalise to `@<semitone>`; nine_of reads them
            # by pitch class, and anything with no image stays unknown
            k = pp.nine_of(ch) if ch is not None else None
            if not k:
                picks.append((ch, '???')); continue
            picks.append((k, rng.choice(po.W[k][po.ROLES[slot]])))
        out.append(po.sentence(picks))
    return ' … then '.join(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--pools', default='G50')
    ap.add_argument('--variants', type=int, default=3)
    ap.add_argument('--out', default=os.path.expanduser('~/Desktop/pao-guitar50.xlsx'))
    ap.add_argument('--seed', type=int, default=11)
    a = ap.parse_args()

    rng = random.Random(a.seed)
    want = set(a.pools.split(','))
    pm = json.load(open(pp.POOLS))
    slugs = sorted(s for s, p in pm.items() if p in want)
    idx = pp.build_index()

    rows, missing = [], []
    for slug in slugs:
        path = pp.find(slug, idx)
        if not path:
            missing.append(slug); continue
        tonic, scale, secs = pp.sections_of(path)
        band, _, title = pp.strip_tags(slug).partition('_')
        for part, seq in secs:
            rows.append(dict(
                song=title.replace('-', ' ').title(),
                band=band.replace('-', ' ').title(),
                part=part,
                key=f'{tonic} {scale}',
                chords=len(seq),
                progression='-'.join(x or '?' for x in seq),
                paos=[scene_for(seq, rng) for _ in range(a.variants)],
            ))

    wb = Workbook(); ws = wb.active; ws.title = 'PAO'
    cols = ([('song', 30), ('band', 24), ('part', 16), ('key', 13),
             ('chords', 8), ('progression', 30)] +
            [(f'pao{i+1}', 62) for i in range(a.variants)])
    ws.append([c[0] for c in cols])
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i); c.fill = HEAD; c.font = HF
    for r in rows:
        ws.append([r['song'], r['band'], r['part'], r['key'], r['chords'],
                   r['progression'], *r['paos']])
    for row in range(2, ws.max_row + 1):
        for col in range(1, len(cols) + 1):
            ws.cell(row, col).alignment = Alignment(vertical='top', wrap_text=col > 6)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{ws.max_row}'
    wb.save(a.out)

    short = sum(1 for r in rows if r['chords'] <= 4)
    print(f'songs      {len(slugs) - len(missing)}/{len(slugs)}')
    print(f'rows       {len(rows)}   ({short} fit one scene, {len(rows)-short} need more)')
    print(f'variants   {a.variants} per row')
    print(f'-> {a.out}')
    if missing:
        print(f'no hookpad file: {", ".join(missing[:6])}')


if __name__ == '__main__':
    main()
