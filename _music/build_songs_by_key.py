"""Regenerate _music/songs-by-key.html from current data.

Sources: keys from UG tab txts (~/Desktop/music/ug_tabs/*.txt 'Key:' line) then
Supabase key_tonic; per-song UG/Hookpad links from Supabase ug_url/hookpad_url +
the guitar_*_ug_urls.txt lists, with a UG-search fallback. Songs = guitar sheets
of songs_to_learn.xlsx.
"""
import openpyxl, os, re, glob, json, csv, urllib.parse
from collections import defaultdict
from dotenv import load_dotenv
load_dotenv('/Users/robert/Desktop/themap/themap_claude/.env')
import psycopg2
def norm(s):
    s=str(s or '').lower(); s=re.sub(r"['']",'',s); s=re.sub(r'[^a-z0-9]+',' ',s); return re.sub(r'\s+',' ',s).strip()
ALI={'creedence clearwater revival':'ccr','red hot chili peppers':'rhcp','lynyrd skynyrd':'lynryd skynyrd','the band':'band','guns n roses':'gnr'}
def na(s):
    n=norm(s); n=n[4:] if n.startswith('the ') else n; return ALI.get(n,n)
def _ug_artist(url):
    m=re.search(r'/tab/([^/]+)/',url or ''); return na(m.group(1).replace('-',' ')) if m else None
def ug_ok(url, artist):
    """True if a UG tab url's artist-slug matches the song artist (or it's a non-tab/search url)."""
    if not url or 'tabs.ultimate-guitar.com' not in url: return True
    ua, sa = _ug_artist(url), na(artist)
    if not ua or not sa: return True
    uc, sc = ua.replace(' ',''), sa.replace(' ','')
    if uc==sc or uc in sc or sc in uc: return True
    return bool({w for w in ua.split() if len(w)>=3} & {w for w in sa.split() if len(w)>=3})
def normkey(raw):
    m=re.match(r'([A-G][#b]?)\s*(m|min|minor)?',(raw or '').strip()); return f"{m.group(1)} {'minor' if m.group(2) else 'major'}" if m else None
LETTER_PC={'C':0,'D':2,'E':4,'F':5,'G':7,'A':9,'B':11}
def tonic_pc(n):
    pc=LETTER_PC[n[0]]; pc+=1 if '#' in n else 0; pc-=1 if 'b' in n[1:] else 0; return pc%12
SEL={0:'C',7:'G',2:'D',9:'A',4:'E',11:'B',6:'F#',1:'Db',8:'Ab',3:'Eb',10:'Bb',5:'F'}

ugkey={}
for f in glob.glob(os.path.expanduser('~/Desktop/music/ug_tabs/*.txt')):
    b=os.path.basename(f)[:-4]; p=b.split('_',1); a,t=(p[0],p[1]) if len(p)==2 else ('',b)
    try: head=open(f,encoding='utf-8',errors='ignore').read(2000)
    except: continue
    mk=re.search(r'Key:\s*([A-G][#b]?m?(?:in)?)',head)
    if mk: ugkey[(na(a),norm(t))]=normkey(mk.group(1))
c=psycopg2.connect(host=os.environ['DB_HOST'],dbname=os.environ['DB_NAME'],user=os.environ['DB_USER'],password=os.environ['DB_PASSWORD'],port=os.environ.get('DB_PORT',5432)); cur=c.cursor()
cur.execute("select artist,title,key_tonic,key_scale,ug_url,hookpad_url from parcels.songs")
sbkey={}; sburl={}
for a,t,kt,ks,ug,hp in cur.fetchall():
    k=(na(a),norm(t))
    if kt: sbkey.setdefault(k,f"{kt} {ks or 'major'}")
    if ug or hp:
        e=sburl.get(k,{}); e['ug']=e.get('ug') or ug; e['hp']=e.get('hp') or hp; sburl[k]=e
c.close()
ugtxt={}
for path in glob.glob(os.path.expanduser('~/Desktop/5-24-26/guitar*_ug_urls.txt')):
    for r in csv.DictReader(open(path),delimiter='\t'):
        k=(na(r.get('artist')),norm(r.get('title')))
        if r.get('url') and k not in ugtxt: ugtxt[k]=r['url']

wb=openpyxl.load_workbook(os.path.expanduser('~/Desktop/music/songs_to_learn.xlsx'),read_only=True,data_only=True)
seen=set(); bykey=defaultdict(list); nokey=0
for name in [s for s in wb.sheetnames if s.lower().replace(' ','').startswith('guitar')]:
    for row in wb[name].iter_rows(values_only=True):
        if not row or all(x is None for x in row): continue
        c2=row[2] if len(row)>2 else None
        if isinstance(c2,str) and '_' in c2 and len(c2.split('_',1))==2: art,tit=c2.split('_',1)
        else: tit,art=row[0],(row[1] if len(row)>1 else None)
        if not tit or norm(tit)=='title': continue
        art=str(art or '').strip(); tit=str(tit).strip(); k=(na(art),norm(tit))
        if k in seen: continue
        seen.add(k)
        key=ugkey.get(k) or sbkey.get(k)
        if not key: nokey+=1; continue
        urls=sburl.get(k,{})
        search='https://www.ultimate-guitar.com/search.php?title='+urllib.parse.quote(tit)+'&artist='+urllib.parse.quote(art)
        # use first candidate whose artist matches the song; else fall back to UG search
        ug=next((u for u in (urls.get('ug'), ugtxt.get(k)) if u and ug_ok(u,art)), search)
        bykey[key].append([tit,art,ug,urls.get('hp')])
def keyobj(k):
    tn=k.split()[0]; minor=k.endswith('minor')
    return {"key":k,"minor":minor,"fret":SEL.get((tonic_pc(tn)+(3 if minor else 0))%12),"songs":[list(s) for s in sorted(set(map(tuple,bykey[k])))]}
majors=sorted([k for k in bykey if k.endswith('major')],key=lambda k:-len(bykey[k]))
minors=sorted([k for k in bykey if not k.endswith('major')],key=lambda k:-len(bykey[k]))
data=[keyobj(k) for k in majors]+[keyobj(k) for k in minors]
total=sum(len(o['songs']) for o in data)
print(f'{len(data)} keys, {total} keyed songs, {nokey} unkeyed')

TMPL=r'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Songs by Key — Glowing Gardens</title><style>
 *{box-sizing:border-box;margin:0;padding:0}
 body{background:#0f0f1f;color:#e0e0e0;font-family:-apple-system,BlinkMacSystemFont,sans-serif;height:100vh;overflow:hidden;display:flex;flex-direction:column}
 header{padding:12px 16px;border-bottom:1px solid #2a2a4a;display:flex;align-items:center;gap:14px}
 header h1{font-size:16px;margin:0} header .tot{margin-left:auto;font-size:13px;color:#8a8ab0;font-family:ui-monospace,monospace}
 a.back{color:#6a6a8a;text-decoration:none;font-size:13px} a.back:hover{color:#fff}
 .wrap{flex:1;display:flex;min-height:0}
 .keys{width:220px;border-right:1px solid #2a2a4a;overflow-y:auto;padding:10px}
 .grp{font-size:11px;text-transform:uppercase;letter-spacing:1px;color:#5a5a7a;margin:12px 8px 6px}
 .keybtn{display:flex;justify-content:space-between;width:100%;text-align:left;background:none;border:none;color:#d0d0e0;padding:8px 10px;border-radius:7px;cursor:pointer;font-size:14px;font-family:ui-monospace,monospace}
 .keybtn:hover{background:#1e1e38} .keybtn.active{background:#2a2a4a;color:#fff;font-weight:700} .keybtn .n{color:#7a7a9a;font-size:12px}
 main{flex:1;overflow-y:auto;padding:18px 24px}
 .khead{display:flex;align-items:baseline;gap:14px;margin-bottom:12px;flex-wrap:wrap}
 .khead h2{font-size:24px;margin:0} .khead .cnt{color:#8a8ab0;font-size:14px}
 .shapes{margin-left:auto;background:#3050d0;color:#fff;text-decoration:none;padding:9px 16px;border-radius:8px;font-size:13px;font-weight:700} .shapes:hover{background:#3b5ae0}
 .filter{margin:0 0 12px;width:100%;background:#1a1a2e;border:1px solid #3a3a5a;color:#e0e0e0;border-radius:7px;padding:8px 10px;font-size:13px}
 .search{background:#1a1a2e;border:1px solid #3a3a5a;color:#e0e0e0;border-radius:7px;padding:7px 11px;font-size:13px;width:210px}
 .song .kb{font-size:10px;color:#9a9ac0;background:#1e1e38;padding:1px 6px;border-radius:5px;margin-left:4px;font-family:ui-monospace,monospace}
 .songs{columns:2;column-gap:28px} @media(max-width:760px){.songs{columns:1}}
 .song{break-inside:avoid;padding:6px 4px;border-bottom:1px solid #1a1a2e;font-size:14px;display:flex;align-items:baseline;gap:6px}
 .song .a{color:#8a8ab0} .song .t{color:#e8e8f4} .song .txt{flex:1}
 .song a.lk{font-size:10px;font-weight:700;text-decoration:none;padding:2px 6px;border-radius:5px}
 a.ug{background:#4a2a2a;color:#ff9b7a} a.ug:hover{background:#5a3232} a.hp{background:#22304a;color:#7ab8ff} a.hp:hover{background:#2a3c5c}
</style></head><body>
<header><a class="back" href="index.html">&larr; Music</a><h1>Songs by Key</h1><input id="search" class="search" placeholder="🔍 search all songs"><span class="tot">__TOTAL__ songs</span></header>
<div class="wrap"><nav class="keys" id="keyList"></nav><main id="pane"></main></div>
<script>
const DATA=__DATA__;
let cur=0;
function render(){const kl=document.getElementById('keyList');
 const mk=(arr,label)=>`<div class="grp">${label}</div>`+arr.map(d=>`<button class="keybtn" data-i="${DATA.indexOf(d)}"><span>${d.key.replace(' major','').replace(' minor','m')}</span><span class="n">${d.songs.length}</span></button>`).join('');
 kl.innerHTML=mk(DATA.filter(d=>!d.minor),'Major')+mk(DATA.filter(d=>d.minor),'Minor');
 kl.querySelectorAll('.keybtn').forEach(b=>b.onclick=()=>select(+b.dataset.i)); select(0);}
function select(i){cur=i;document.querySelectorAll('.keybtn').forEach(b=>b.classList.toggle('active',+b.dataset.i===i));
 const d=DATA[i];
 const shapes=d.fret?`<a class="shapes" href="../guitar.html?scale=${encodeURIComponent(d.fret)}">🎸 ${d.fret} major scale${d.minor?' (relative)':''}</a>`:'';
 const rows=d.songs.map(([t,a,ug,hp])=>`<div class="song"><span class="txt"><span class="t">${esc(t)}</span> <span class="a">— ${esc(a)}</span></span>${ug?`<a class="lk ug" target="_blank" href="${esc(ug)}">UG</a>`:''}${hp?`<a class="lk hp" target="_blank" href="${esc(hp)}">HP</a>`:''}</div>`).join('');
 document.getElementById('pane').innerHTML=`<div class="khead"><h2>${d.key}</h2><span class="cnt">${d.songs.length} songs</span>${shapes}</div><input class="filter" placeholder="filter…" oninput="flt(this.value)"><div class="songs" id="songs">${rows}</div>`;}
function flt(q){q=q.toLowerCase();document.querySelectorAll('#songs .song').forEach(s=>{s.style.display=s.textContent.toLowerCase().includes(q)?'':'none';});}
function esc(s){return (s||'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));}
function searchAll(q){const ql=q.toLowerCase(),hits=[];
 DATA.forEach(d=>d.songs.forEach(s=>{if((s[0]+' '+s[1]).toLowerCase().includes(ql))hits.push(s.concat([d.key]));}));
 document.querySelectorAll('.keybtn').forEach(b=>b.classList.remove('active'));
 const rows=hits.map(([t,a,ug,hp,key])=>`<div class="song"><span class="txt"><span class="t">${esc(t)}</span> <span class="a">— ${esc(a)}</span> <span class="kb">${key}</span></span>${ug?`<a class="lk ug" target="_blank" href="${esc(ug)}">UG</a>`:''}${hp?`<a class="lk hp" target="_blank" href="${esc(hp)}">HP</a>`:''}</div>`).join('');
 document.getElementById('pane').innerHTML=`<div class="khead"><h2>Search</h2><span class="cnt">${hits.length} matches</span></div><div class="songs">${rows||'<div style="color:#8a8ab0">no matches</div>'}</div>`;}
document.getElementById('search').oninput=e=>{const q=e.target.value.trim();q?searchAll(q):select(cur);};
render();
</script></body></html>'''
open('/Users/robert/Desktop/glowinggardens_claude/_music/songs-by-key.html','w').write(
    TMPL.replace('__TOTAL__',str(total)).replace('__DATA__',json.dumps(data,ensure_ascii=False)))
print('wrote _music/songs-by-key.html')
